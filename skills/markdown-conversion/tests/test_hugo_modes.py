"""Standalone CLI and image-mode regressions for Hugo adaptations."""

import base64
import contextlib
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

import convert
import doc_to_md
import pdf_to_md
import ppt_to_md
import web_to_md
from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import fitz


class ImageModeTests(unittest.TestCase):
    def _fixtures(self, root):
        large = root / "figure.png"
        small = root / "icon.png"
        Image.new("RGB", (400, 200), "white").save(large)
        Image.new("RGB", (20, 20), "white").save(small)

        doc = Document()
        doc.add_paragraph("KEEP_BODY")
        doc.add_picture(str(large))
        doc.add_picture(str(small))
        doc.save(root / "sample.docx")

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "KEEP_BODY"
        slide.shapes.add_picture(str(large), Inches(1), Inches(2), width=Inches(4))
        slide.shapes.add_picture(str(small), Inches(1), Inches(5), width=Inches(.2))
        deck.save(root / "sample.pptx")

        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "KEEP_BODY")
        page.insert_image(fitz.Rect(72, 100, 472, 300), filename=str(large))
        page.insert_image(fitz.Rect(72, 400, 92, 420), filename=str(small))
        pdf.save(root / "sample.pdf")
        pdf.close()

        images = "".join('<img src="data:image/png;base64,' +
                         base64.b64encode(p.read_bytes()).decode() + '">' for p in (large, small))
        (root / "sample.html").write_text("<html><body><p>KEEP_BODY</p>" + images +
                                         "</body></html>", encoding="utf-8")

    def test_native_backends_keep_filter_or_skip_images(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self._fixtures(root)
            for kind in ("docx", "pptx", "pdf", "html"):
                for mode, expected in (("all", 2), ("filtered", 1), ("none", 0)):
                    with self.subTest(kind=kind, mode=mode):
                        output = root / f"{kind}-{mode}.md"
                        options = dict(no_images=mode == "none", filter_images=mode == "filtered")
                        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                            if kind == "pdf":
                                text = pdf_to_md.extract_pdf_to_markdown(
                                    str(root / "sample.pdf"), str(output), images=mode)
                            elif kind == "pptx":
                                text = ppt_to_md.convert_presentation_to_markdown(
                                    str(root / "sample.pptx"), str(output), **options)
                            else:
                                text = doc_to_md.convert_to_markdown(
                                    str(root / f"sample.{kind}"), str(output), **options)
                        self.assertIn("KEEP_BODY", text.replace(chr(92), ""))
                        self.assertEqual(text.count("!["), expected)
                        assets = output.parent / (output.stem + "_files")
                        self.assertEqual(len(list(assets.glob("*.png"))), expected)
                        profile = json.loads(output.with_suffix(".conversion_profile.json").read_text())
                        self.assertEqual(profile["schema"], "markdown-conversion.conversion_profile.v1")
                        self.assertEqual(profile["outputs"]["image_count"], expected)

    def test_web_modes_keep_filter_or_skip_images_and_provenance(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self._fixtures(root)
            page = b'<html><body><p>KEEP_BODY</p><img src="figure.png"><img src="icon.png"></body></html>'
            def response(url, **kwargs):
                body = page if url.endswith("/page") else (root / Path(url).name).read_bytes()
                return SimpleNamespace(url=url, content=body, headers={"Content-Type":
                    "text/html; charset=utf-8" if url.endswith("/page") else "image/png"},
                    raise_for_status=lambda: None)
            for mode, expected in (("all", 2), ("filtered", 1), ("none", 0)):
                output = root / f"web-{mode}.md"
                with patch.object(web_to_md, "fetch_response", side_effect=response),                      patch.object(web_to_md, "_http_get", side_effect=response),                      contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    result = web_to_md.process_url(
                        "https://example.invalid/page", str(output), raw=True,
                        no_images=mode == "none", filter_images=mode == "filtered")
                self.assertTrue(result[0], result)
                self.assertEqual(output.read_text().count("!["), expected)
                if expected:
                    sources = json.loads((root / (output.stem + "_files") / "image_sources.json").read_text())
                    self.assertEqual(len(sources["items"]), expected)
                    self.assertTrue(all(item["license_status"] == "unknown" and
                        item["download_url"] and item["local_path"] for item in sources["items"]))


    def test_epub_and_notebook_image_modes(self):
        import ebooklib
        from ebooklib import epub
        import nbformat

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self._fixtures(root)
            encoded = base64.b64encode((root / "figure.png").read_bytes()).decode()
            book = epub.EpubBook()
            book.set_identifier("fixture")
            book.set_title("Fixture")
            book.set_language("en")
            chapter = epub.EpubHtml(title="Body", file_name="body.xhtml")
            chapter.content = '<html><body><p>Keep body</p><img src="figure.png"/></body></html>'
            book.add_item(chapter)
            book.add_item(epub.EpubItem(uid="figure", file_name="figure.png",
                                       media_type="image/png", content=(root / "figure.png").read_bytes()))
            book.toc = [chapter]
            book.add_item(epub.EpubNcx())
            book.add_item(epub.EpubNav())
            book.spine = [chapter]
            epub.write_epub(str(root / "sample.epub"), book)
            cell = nbformat.v4.new_markdown_cell("Keep body\n\n![figure](attachment:figure.png)")
            cell.attachments = {"figure.png": {"image/png": encoded}}
            notebook = nbformat.v4.new_notebook(cells=[cell])
            nbformat.write(notebook, root / "sample.ipynb")
            for kind in ("epub", "ipynb"):
                for mode in ("all", "filtered", "none"):
                    with self.subTest(kind=kind, mode=mode), \
                         contextlib.redirect_stdout(io.StringIO()), \
                         contextlib.redirect_stderr(io.StringIO()):
                        output = root / f"{kind}-{mode}.md"
                        text = doc_to_md.convert_to_markdown(
                            str(root / f"sample.{kind}"), str(output),
                            no_images=mode == "none", filter_images=mode == "filtered")
                        self.assertIn("Keep body", text)
                        self.assertEqual(text.count("!["), 0 if mode == "none" else 1)

    def test_pdf_no_images_disables_vector_rendering(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self._fixtures(root)
            with patch.object(pdf_to_md, "detect_vector_figure_rects") as detect, \
                 contextlib.redirect_stdout(io.StringIO()):
                text = pdf_to_md.extract_pdf_to_markdown(
                    str(root / "sample.pdf"), str(root / "vector.md"),
                    images="none", render_vector_figures=True, raw=True)
            detect.assert_not_called()
            self.assertIn("KEEP_BODY", text)
            self.assertNotIn("![", text)

    def test_pptx_soft_breaks_survive(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "alpha\vbeta"
            source = root / "breaks.pptx"
            deck.save(source)
            with contextlib.redirect_stdout(io.StringIO()):
                text = ppt_to_md.convert_presentation_to_markdown(str(source), str(root / "breaks.md"))
            self.assertIn("alpha\nbeta", text)


class UnifiedEntryTests(unittest.TestCase):
    def test_dispatcher_forwards_insecure_to_web_backend(self):
        with patch.object(convert, "run_python_script", return_value=1) as run:
            result = convert.main(["https://example.org/page", "-o", "/unused/result.md", "--insecure"])
        self.assertEqual(result, 1)
        script_name, command = run.call_args.args
        self.assertEqual(script_name, "web_to_md.py")
        self.assertEqual(command.count("--insecure"), 1)

    def test_office_success_emits_output_and_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            doc = Document()
            doc.add_paragraph("KEEP_BODY")
            source = root / "source.docx"
            doc.save(source)
            for mode in ([], ["--filter-images"], ["--no-images"]):
                output = root / ("out" + str(len(mode)) + ("none" if "--no-images" in mode else "") + ".md")
                result = subprocess.run([sys.executable, str(SCRIPTS_DIR / "convert.py"),
                    str(source), "-o", str(output), "--json", *mode],
                    capture_output=True, text=True, timeout=60)
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertIn("OUTPUT: " + str(output.resolve()), result.stdout)
                payload = json.loads(result.stdout.splitlines()[-1])
                self.assertEqual(payload["markdown"], str(output.resolve()))
                self.assertTrue(Path(payload["conversion_profile"]).is_file())

    def test_mutually_exclusive_modes_are_rejected(self):
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(convert.main(["sample.pdf", "--no-images", "--filter-images"]), 2)

    def test_web_downloaded_document_receives_hugo_modes(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "sample.md"
            def run(command):
                output.write_text("KEEP_BODY")
                self.assertIn("--no-images", command)
                self.assertIn("--raw", command)
                return SimpleNamespace(returncode=0)
            with patch.object(web_to_md.subprocess, "run", side_effect=run),                  contextlib.redirect_stdout(io.StringIO()):
                result = web_to_md._convert_remote_document(
                    "https://example.invalid/source.docx", b"fixture", ".docx", str(output),
                    no_images=True, raw=True)
            self.assertTrue(result[0], result)

    def test_trafilatura_preserves_links_and_falls_back_when_empty(self):
        html = b'<html><body><p>KEEP_BODY <a href="/proof">Proof</a></p></body></html>'
        response = SimpleNamespace(url="https://example.invalid/page", content=html,
                                   headers={"Content-Type": "text/html; charset=utf-8"})
        for extracted in ('<p>KEEP_BODY <a href="/proof">Proof</a></p>', None):
            extractor = SimpleNamespace(extract=lambda *args, **kwargs: extracted)
            with tempfile.TemporaryDirectory() as tmp,                  patch.object(web_to_md, "trafilatura", extractor),                  patch.object(web_to_md, "fetch_response", return_value=response),                  contextlib.redirect_stdout(io.StringIO()):
                output = Path(tmp) / "page.md"
                result = web_to_md.process_url(response.url, str(output), no_images=True)
                self.assertTrue(result[0], result)
                self.assertIn("[Proof](https://example.invalid/proof)", output.read_text())


if __name__ == "__main__":
    unittest.main()
