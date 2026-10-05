from __future__ import annotations
import sys
from pathlib import Path
SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))
import json

import sys

import tempfile

import unittest

import zipfile

from pathlib import Path

from xml.etree import ElementTree as ET

import pdf_to_md

import doc_to_md

import ppt_to_md

import web_to_md

from _conversion_profile import profile_path_for, record_source_url, write_conversion_profile

class PresentationSoftBreakTests(unittest.TestCase):
    def test_soft_line_break_survives_readback(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "Section one\vOverall duties"  # \v writes <a:br/>
        markdown = ppt_to_md.text_frame_to_markdown(box.text_frame, box)
        self.assertIn("Section one\nOverall duties", markdown)


class RemoteDocumentSuffixTests(unittest.TestCase):
    def test_web_to_md_refuses_output_file_for_several_urls(self) -> None:
        import io
        from contextlib import redirect_stderr
        buffer = io.StringIO()
        with redirect_stderr(buffer):
            rc = web_to_md.main(["https://example.org/a", "https://example.org/b", "-o", "out.md"])
        self.assertEqual(rc, 2)
        self.assertIn("--dir", buffer.getvalue())

    def test_pdf_url_with_pdf_content_type(self) -> None:
        self.assertEqual(
            web_to_md.remote_document_suffix(PDF_URL, "application/pdf", b"%PDF-1.4"), ".pdf")

    def test_suffixless_download_answering_pdf(self) -> None:
        self.assertEqual(
            web_to_md.remote_document_suffix(
                "https://example.org/download?id=7", "application/pdf; charset=binary", b"%PDF-1.7"),
            ".pdf")

    def test_body_magic_wins_over_wrong_content_type(self) -> None:
        self.assertEqual(
            web_to_md.remote_document_suffix(
                "https://example.org/files/report", "application/octet-stream", b"%PDF-1.7"),
            ".pdf")

    def test_docx_url_by_suffix(self) -> None:
        self.assertEqual(
            web_to_md.remote_document_suffix(
                "https://example.org/a/brief.docx", "application/octet-stream", b"PK\x03\x04"),
            ".docx")

    def test_html_viewer_at_document_url_stays_web(self) -> None:
        self.assertIsNone(
            web_to_md.remote_document_suffix(PDF_URL, "text/html; charset=utf-8", b"<!doctyp"))

    def test_ordinary_page_is_not_a_document(self) -> None:
        self.assertIsNone(
            web_to_md.remote_document_suffix(
                "https://example.org/article", "text/html", b"<html>"))
        self.assertIsNone(
            web_to_md.remote_document_suffix(
                "https://example.org/page.html", "", b"<html>"))

class PdfBulletTests(unittest.TestCase):
    def test_middot_bullet_is_a_list_item(self) -> None:
        is_list, kind, content = pdf_to_md.detect_list_item("·\x01 Oaks turn red, brown, or russet;")
        self.assertTrue(is_list)
        self.assertEqual(kind, "ul")
        self.assertEqual(content, "- Oaks turn red, brown, or russet;")

    def test_bullet_led_line_is_never_a_heading(self) -> None:
        size_map = {"body": 12.1, "h1": 13.6}
        self.assertEqual(pdf_to_md.get_heading_level(13.6, size_map, "· Oaks turn red", 4), 0)
        self.assertEqual(pdf_to_md.get_heading_level(13.6, size_map, "Autumn colours", 16), 1)

    def test_bullet_glyph_span_detection(self) -> None:
        self.assertTrue(pdf_to_md.is_bullet_glyph_span("· "))
        self.assertTrue(pdf_to_md.is_bullet_glyph_span("•"))
        self.assertTrue(pdf_to_md.is_bullet_glyph_span("·\x01"))
        self.assertFalse(pdf_to_md.is_bullet_glyph_span("Oaks"))
        self.assertFalse(pdf_to_md.is_bullet_glyph_span("· Oaks"))

class PdfClauseAndScanTests(unittest.TestCase):
    def test_spaced_clause_number_is_not_a_list(self) -> None:
        self.assertFalse(pdf_to_md.detect_list_item("1. 1 职业名称")[0])
        self.assertFalse(pdf_to_md.detect_list_item("2. 1. 1 职业道德基本知识")[0])

    def test_ordinary_ordered_items_still_match(self) -> None:
        self.assertEqual(pdf_to_md.detect_list_item("1. 职业概况"), (True, "ol", "1. 职业概况"))
        self.assertEqual(pdf_to_md.detect_list_item("3. Overview"), (True, "ol", "3. Overview"))
        self.assertEqual(pdf_to_md.detect_list_item("1. 2024年营收"), (True, "ol", "1. 2024年营收"))
        self.assertFalse(pdf_to_md.detect_list_item("83.2% of respondents")[0])

    def test_scanned_pdf_warns(self) -> None:
        markdown = "\n".join(
            f"<!-- Page {i} -->\n![page {i}](scan_files/page_{i}.jpg)" for i in range(1, 14))
        warnings = pdf_to_md.scanned_pdf_warnings(markdown, 13, 13)
        self.assertEqual(len(warnings), 1)
        self.assertIn("13 page images", warnings[0])

    def test_text_pdf_does_not_warn(self) -> None:
        markdown = "\n".join("第一章 总则 " * 20 for _ in range(13))
        self.assertEqual(pdf_to_md.scanned_pdf_warnings(markdown, 13, 2), [])

class JapaneseTypographyTests(unittest.TestCase):
    def test_wrapped_pdf_lines_join_without_space_between_cjk(self) -> None:
        self.assertEqual(pdf_to_md.join_wrapped_text("約６８億", "人ものお客様"), "約６８億人ものお客様")
        self.assertEqual(pdf_to_md.join_wrapped_text("新幹", "線"), "新幹線")
        self.assertEqual(pdf_to_md.join_wrapped_text("the high", "speed"), "the high speed")

class TraditionalChineseIntakeTests(unittest.TestCase):
    def test_table_cells_keep_links(self) -> None:
        html = (
            '<table><tr><td>報告</td><td><ul><li><a href="/a.pdf">pdf</a></li>'
            '<li><a href="/a.docx">docx</a></li></ul></td></tr></table>'
        )
        markdown = web_to_md.simple_html_to_markdown_traversal(
            web_to_md.BeautifulSoup(html, "html.parser"), "https://example.gov.tw/x",
        )
        self.assertIn("[pdf](https://example.gov.tw/a.pdf)", markdown)
        self.assertIn("[docx](https://example.gov.tw/a.docx)", markdown)
    def test_downloaded_document_profile_records_url(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            markdown = Path(tmp) / "report.md"
            markdown.write_text("# r\n", encoding="utf-8")
            write_conversion_profile(
                input_path=str(Path(tmp) / "report.pdf"), markdown_path=markdown,
                converter="pdf_to_md.py", conversion_type="pdf",
            )
            record_source_url(markdown, "https://example.gov.tw/report.pdf")
            profile = json.loads(profile_path_for(markdown).read_text(encoding="utf-8"))
        self.assertEqual(profile["source"]["url"], "https://example.gov.tw/report.pdf")

class DocxIntakeTests(unittest.TestCase):
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    DOCUMENT = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><w:body>'
        '<w:p><w:r><w:drawing><c:chart r:id="rId9"/></w:drawing></w:r></w:p>'
        '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Bar queues</w:t></w:r>'
        '<w:r><w:footnoteReference w:id="2"/></w:r></w:p></w:tc></w:tr></w:tbl>'
        '</w:body></w:document>'
    )
    RELS = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId9" Type="chart" Target="charts/chart1.xml"/></Relationships>'
    )
    CHART = (
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart><c:plotArea><c:barChart><c:barDir val="col"/><c:grouping val="clustered"/>'
        '<c:ser><c:idx val="0"/><c:order val="0"/>'
        '<c:cat><c:strRef><c:strCache><c:ptCount val="2"/>'
        '<c:pt idx="0"><c:v>2022/23</c:v></c:pt><c:pt idx="1"><c:v>2023/24</c:v></c:pt>'
        '</c:strCache></c:strRef></c:cat>'
        '<c:val><c:numRef><c:numCache><c:ptCount val="2"/>'
        '<c:pt idx="0"><c:v>702</c:v></c:pt><c:pt idx="1"><c:v>659</c:v></c:pt>'
        '</c:numCache></c:numRef></c:val></c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>'
    )
    FOOTNOTES = (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
        '<w:footnote w:id="2"><w:p><w:r><w:t>Theatre closed until Winter 2026.</w:t></w:r></w:p>'
        '</w:footnote></w:footnotes>'
    )

    def _docx(self, root: Path) -> Path:
        path = root / "report.docx"
        with zipfile.ZipFile(path, "w") as docx:
            docx.writestr("word/document.xml", self.DOCUMENT)
            docx.writestr("word/_rels/document.xml.rels", self.RELS)
            docx.writestr("word/charts/chart1.xml", self.CHART)
            docx.writestr("word/footnotes.xml", self.FOOTNOTES)
        return path

    def test_embedded_chart_cache_becomes_a_table(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            patched, replacements, warnings = doc_to_md._docx_inject_charts_markdown(
                self._docx(Path(tmp)))
            patched.unlink()
        [markdown] = replacements.values()
        self.assertIn("| 2023/24 | 659 |", markdown)
        self.assertEqual(warnings, [])

    def test_table_cell_footnote_survives(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            patched, replacements = doc_to_md._docx_inject_tables_markdown(self._docx(Path(tmp)))
            patched.unlink()
        [markdown] = replacements.values()
        self.assertIn("Bar queues[^2]", markdown)
        self.assertIn("[^2]: Theatre closed until Winter 2026.", markdown)

class KoreanIntakeTests(unittest.TestCase):
    def test_pdf_join_keeps_korean_word_space(self) -> None:
        self.assertEqual(pdf_to_md.join_wrapped_text("감소하였으며", "이중"), "감소하였으며 이중")
        self.assertEqual(pdf_to_md.join_wrapped_text("新幹", "線"), "新幹線")
    def test_web_table_is_gfm_on_its_grid(self) -> None:
        html = (
            "<table><caption>해녀 현황</caption>"
            '<tr><th rowspan="2">구분</th><th colspan="2">계</th></tr>'
            "<tr><th>2025</th><th>2024</th></tr>"
            "<tr><td>계</td><td>7,482</td><td>7,561</td></tr></table>"
        )
        markdown = web_to_md.simple_html_to_markdown_traversal(
            web_to_md.BeautifulSoup(html, "html.parser"), "https://example.kr/")
        self.assertIn("해녀 현황\n\n| 구분 | 계 |  |\n| --- | --- | --- |", markdown)
        self.assertIn("|  | 2025 | 2024 |\n| 계 | 7,482 | 7,561 |", markdown)

class ArabicPdfTests(unittest.TestCase):
    def test_reversed_lam_alef_layer_warns(self) -> None:
        broken = "نشرة اإلحصاءات الزراعية األعلى آالف " * 40
        self.assertEqual(len(pdf_to_md.arabic_text_layer_warnings(broken)), 1)

    def test_well_formed_arabic_does_not_warn(self) -> None:
        clean = "القهوة العربية رمز للكرم والضيافة في شبه الجزيرة العربية " * 20
        self.assertEqual(pdf_to_md.arabic_text_layer_warnings(clean), [])

class IntakeHousekeepingTests(unittest.TestCase):
    def test_record_numbers_in_urls_are_not_dates(self) -> None:
        from bs4 import BeautifulSoup
        empty = BeautifulSoup("<html><title>x</title></html>", "html.parser")
        date = lambda url: web_to_md.extract_metadata(empty, url)["date"]
        self.assertEqual(date("https://iris.who.int/bitstream/handle/10665/379812/x.pdf"), "")
        self.assertEqual(date("https://www.mem.gov.cn/kp/shaq/202205/t20220519_413952.shtml"), "2022-05")

PDF_URL = "https://www.example.gov/content/pkg/report/pdf/report.pdf"
